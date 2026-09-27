#!/usr/bin/env python3
"""Validate the published series against external benchmarks.

    python scripts/validate.py ln2015
        Value-weighted unlevered Core index vs the Ling & Naranjo (2015) unlevered core
        REIT index (column "L&N unlev (core)" of legacy/2021-04/Unlevered Return Indices.xlsx),
        1993Q1-2012Q4.

    python scripts/validate.py nareit --nareit-file MonthlyHistoricalReturns.xls
        Value-weighted levered All-REIT index vs the FTSE Nareit All Equity REITs total
        return. Download the monthly history workbook from https://www.reit.com (FTSE
        Nareit U.S. Real Estate Index Series); it is not redistributed here. Monthly
        returns are compounded to calendar quarters (complete quarters only).
"""
import argparse
import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
INDICES = os.path.join(ROOT, "data", "reit_return_indices.csv")
LEGACY = os.path.join(ROOT, "legacy", "2021-04", "Unlevered Return Indices.xlsx")


def ann(s):
    """Annualized geometric mean of quarterly % returns."""
    return ((1 + s / 100).prod() ** (4 / len(s)) - 1) * 100


def cum(s):
    return ((1 + s / 100).prod() - 1) * 100


def compare(ours, bench, label):
    m = pd.concat([ours.rename("ours"), bench.rename("bench")], axis=1, join="inner").dropna()
    d = m.ours - m.bench
    first, last = m.index.min(), m.index.max()
    print(f"{label}: {len(m)} quarters, {first[0]}Q{first[1]}-{last[0]}Q{last[1]}")
    print(f"  correlation                 {m.ours.corr(m.bench):.4f}")
    print(f"  mean, % per quarter         {m.ours.mean():.3f} vs {m.bench.mean():.3f}"
          f"  (difference {d.mean():+.3f}; mean absolute difference {d.abs().mean():.3f})")
    print(f"  annualized return           {ann(m.ours):.2f} vs {ann(m.bench):.2f}")
    print(f"  annualized volatility       {m.ours.std() * 2:.2f} vs {m.bench.std() * 2:.2f}")
    print(f"  tracking error (annualized) {d.std() * 2:.2f}")
    print(f"  cumulative return           {cum(m.ours):.1f}% vs {cum(m.bench):.1f}%")
    worst = d.reindex(d.abs().sort_values(ascending=False).index).head(3)
    print("  largest quarterly gaps      " + ", ".join(f"{y}Q{q} {v:+.2f}" for (y, q), v in worst.items()))


def load_ours(col):
    df = pd.read_csv(INDICES)
    return df.set_index(["year", "quarter"])[col]


def ln2015(_args):
    ws = load_workbook(LEGACY, read_only=True, data_only=True)["UnlevRet"]
    rows = [r for r in ws.iter_rows(min_row=10, max_col=3, values_only=True)
            if r[0] is not None and r[2] is not None]
    ln = pd.Series({(int(y), int(q)): float(v) for y, q, v in rows})
    compare(load_ours("vw_unlev_core"), ln, "VW unlevered Core vs Ling & Naranjo (2015)")


def nareit(args):
    raw = pd.read_excel(args.nareit_file, sheet_name="Index Data", header=None)
    recs = [(pd.Timestamp(r[0]), pd.to_numeric(r[args.column], errors="coerce"))
            for r in raw.itertuples(index=False) if hasattr(r[0], "year") and not pd.isna(r[0])]
    m = pd.DataFrame(recs, columns=["date", "ret"]).dropna()
    m["year"], m["quarter"] = m.date.dt.year, m.date.dt.quarter
    g = m.groupby(["year", "quarter"]).ret
    q = ((1 + m.ret / 100).groupby([m.year, m.quarter]).prod() - 1) * 100
    q = q[g.size() == 3]
    compare(load_ours("vw_lev_all"), q, "VW levered All vs FTSE Nareit All Equity REITs")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("ln2015", help="compare with Ling & Naranjo (2015)").set_defaults(func=ln2015)
    p = sub.add_parser("nareit", help="compare with FTSE Nareit All Equity REITs")
    p.add_argument("--nareit-file", required=True, help="FTSE Nareit monthly history workbook (.xls)")
    p.add_argument("--column", type=int, default=22,
                   help="0-based column of the All Equity REITs total return on sheet 'Index Data' (default 22)")
    p.set_defaults(func=nareit)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    np.seterr(all="ignore")
    main()
