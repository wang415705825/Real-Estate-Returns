#!/usr/bin/env python3
"""Build every derived output from the published CSVs (public data only):

    data/summary_statistics.csv      annualized statistics by group, weighting, series, period
    data/REIT_Return_Indices.xlsx    Excel copy of the CSVs with a notes sheet
    figures/*.png                    README figures (light and dark) and link-preview cards
    README.md                        the statistics table between the stats markers

Run from anywhere:  python scripts/build_outputs.py
Re-running on unchanged CSVs reproduces byte-identical files.
"""
import glob
import io
import os
import re
import zipfile
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import to_rgba
from matplotlib.patches import FancyBboxPatch, Rectangle
from matplotlib.ticker import FixedLocator, FuncFormatter, MultipleLocator
from matplotlib.transforms import blended_transform_factory
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DATA, FIGS = os.path.join(ROOT, "data"), os.path.join(ROOT, "figures")
FONT_DIR = os.path.join(ROOT, "scripts", "fonts")      # Inter, SIL Open Font License 1.1
LEGACY = os.path.join(ROOT, "legacy", "2021-04", "Unlevered Return Indices.xlsx")
VINTAGE = "2026-09"
STAMP = "2026-09-01T00:00:00Z"          # fixed file timestamp for reproducible xlsx builds

GROUPS = [("All", "all", "All equity REITs"), ("Core", "core", "Core"), ("Non-core", "noncore", "Non-core")]
WEIGHTS = [("VW", "vw", "value-weighted"), ("EW", "ew", "equal-weighted")]
SERIES = [("unlevered", "unlev"), ("levered", "lev")]
PERIODS = [((1993, 1), (2025, 4)), ((1993, 1), (2012, 4)), ((2013, 1), (2025, 4))]

# Chart tokens: reference palette (slot 1 blue = unlevered, slot 2 orange = levered),
# validated for both modes with the dataviz palette checker.
THEMES = {
    "light": dict(surface="#fcfcfb", ink="#0b0b0b", ink2="#52514e", muted="#898781",
                  grid="#e1e0d9", axis="#c3c2b7", band="#f0efec", unlev="#2a78d6", lev="#eb6834"),
    "dark": dict(surface="#1a1a19", ink="#ffffff", ink2="#c3c2b7", muted="#898781",
                 grid="#2c2c2a", axis="#383835", band="#262624", unlev="#3987e5", lev="#d95926"),
}
# Link-preview cards use the author's website colors around the same validated series colors.
CARD = dict(surface="#f7f2e8", panel="#fffdf8", ink="#252422", ink2="#44413c", muted="#6d675e",
            accent="#78283b", rule="#bd9349", grid="#e9e2d6", axis="#d8cec0",
            unlev="#2a78d6", lev="#eb6834")
REPO = "github.com/wang415705825/Real-Estate-Returns"
SOURCE = f"Data: Real-Estate-Returns ({REPO}). Method: Ling and Naranjo (2015)."
# NBER-dated U.S. recessions in the sample, from the start of the peak month to the end of
# the trough month, on the chart's year axis (1993.0 = end of 1992Q4).
RECESSIONS = [(2001 + 2 / 12, 2001 + 11 / 12), (2007 + 11 / 12, 2009 + 6 / 12), (2020 + 1 / 12, 2020 + 4 / 12)]


# ----------------------------------------------------------------------------- data
def load():
    idx = pd.read_csv(os.path.join(DATA, "reit_return_indices.csv"))
    bytype = pd.read_csv(os.path.join(DATA, "reit_returns_by_property_type.csv"))
    return idx, bytype


def qlabel(yq):
    return f"{yq[0]}Q{yq[1]}"


def in_period(df, p):
    k = df.year * 4 + df.quarter
    return df[(k >= p[0][0] * 4 + p[0][1]) & (k <= p[1][0] * 4 + p[1][1])]


def stats(s):
    s = s.dropna()
    if s.empty:
        return None
    return dict(n_quarters=len(s), mean_qtr=s.mean(), sd_qtr=s.std(),
                ann_return=((1 + s / 100).prod() ** (4 / len(s)) - 1) * 100,
                ann_volatility=s.std() * 2, min_qtr=s.min(), max_qtr=s.max())


def summary(idx, bytype):
    rows = []
    for p in PERIODS:
        full = len(in_period(idx, p))
        sub = in_period(idx, p)
        for g, tag, _ in GROUPS:
            for w, wtag, _ in WEIGHTS:
                for sname, stag in SERIES:
                    st = stats(sub[f"{wtag}_{stag}_{tag}"])
                    if st:
                        rows.append(dict(group=g, weighting=w, series=sname, period_start=qlabel(p[0]),
                                         period_end=qlabel(p[1]), **st))
        # property types: only where every quarter of the period is published
        for ty, d in bytype.groupby("property_type"):
            d = in_period(d, p)
            if len(d) != full or d[["vw_unlev", "ew_unlev", "vw_lev", "ew_lev"]].isna().any().any():
                continue
            for w, wtag, _ in WEIGHTS:
                for sname, stag in SERIES:
                    rows.append(dict(group=ty, weighting=w, series=sname, period_start=qlabel(p[0]),
                                     period_end=qlabel(p[1]), **stats(d[f"{wtag}_{stag}"])))
    out = pd.DataFrame(rows)
    num = ["mean_qtr", "sd_qtr", "ann_return", "ann_volatility", "min_qtr", "max_qtr"]
    out[num] = out[num].round(6)
    out.to_csv(os.path.join(DATA, "summary_statistics.csv"), index=False, float_format="%.6f", lineterminator="\n")
    return out


# ----------------------------------------------------------------------------- README table
def readme_table(sm, idx):
    full = sm[(sm.period_start == "1993Q1") & (sm.period_end == "2025Q4")]
    nfull = len(idx)
    lines = ["| Series, 1993Q1–2025Q4 | Unlevered: return | Unlevered: volatility | Levered: return | Levered: volatility |",
             "|---|---:|---:|---:|---:|"]
    flagged = False
    for g, _, gname in GROUPS:
        for w, _, wname in WEIGHTS:
            cells = []
            for sname, _ in SERIES:
                r = full[(full.group == g) & (full.weighting == w) & (full.series == sname)].iloc[0]
                mark = "†" if r.n_quarters < nfull else ""
                flagged |= bool(mark)
                cells += [f"{r.ann_return:.2f}%{mark}", f"{r.ann_volatility:.2f}%{mark}"]
            lines.append(f"| {gname}, {wname} | " + " | ".join(cells) + " |")
    lines.append("")
    lines.append("Annualized geometric mean return and annualized volatility (quarterly standard deviation × 2), "
                 "computed from the published CSV; more periods in [`data/summary_statistics.csv`](data/summary_statistics.csv).")
    if flagged:
        lines.append("")
        lines.append("† Computed over the available quarters: the 1993Q1 value of this series is not yet "
                     "published (see [CHANGELOG.md](CHANGELOG.md)).")
    return "\n".join(lines)


def update_readme(table):
    path = os.path.join(ROOT, "README.md")
    if not os.path.exists(path):
        print(table)
        return
    text = open(path, encoding="utf-8").read()
    new = re.sub(r"(<!-- stats:start -->\n).*?(<!-- stats:end -->)",
                 lambda m: m.group(1) + table + "\n" + m.group(2), text, flags=re.S)
    if new != text:
        open(path, "w", encoding="utf-8", newline="\n").write(new)


# ----------------------------------------------------------------------------- workbook
NOTES = [
    ("Unlevered and levered U.S. equity REIT return indices", True),
    (f"Vintage {VINTAGE}. Quarterly, 1993Q1–2025Q4. https://github.com/wang415705825/Real-Estate-Returns", False),
    ("", False),
    ("Contents", True),
    ("Indices: all equity REITs, Core and Non-core; value-weighted (vw) and equal-weighted (ew); "
     "unlevered (unlev) and levered (lev) total returns; n = number of REITs.", False),
    ("By property type: the same four series for each S&P Global property type. Returns are left blank "
     "when fewer than 3 REITs are available (n_firms is always reported).", False),
    ("Summary statistics: annualized geometric returns and volatilities (quarterly SD x 2).", False),
    ("", False),
    ("Definitions", True),
    ("Returns are nominal total returns in percent per calendar quarter (not annualized).", False),
    ("Levered return: the REIT's CRSP total stock return, including dividends and delisting returns.", False),
    ("Unlevered return: the weighted average of the returns to equity, debt and preferred stock, with weights "
     "equal to each claim's share of total capital at the end of the prior quarter "
     "(Ling and Naranjo, 2015, Real Estate Economics).", False),
    ("Value weights: market capitalization at the end of the prior quarter.", False),
    ("Core: Office, Multifamily, Industrial, Shopping Center, Regional Mall and Other Retail REITs. "
     "Non-core: all other property types.", False),
    ("Universe: equity REITs in S&P Global Market Intelligence with U.S. REIT status, matched to CRSP by CUSIP.", False),
    ("", False),
    ("Sources", True),
    ("Compiled from CRSP (CIZ format) and Compustat via WRDS, and S&P Global Market Intelligence, "
     "retrieved in mid-2026. No vendor data are included; only aggregate index returns.", False),
    ("", False),
    ("Citation", True),
    ("Ling, D. C., Wang, C., and Zhou, T. (2022). Asset productivity, local information diffusion, and "
     "commercial real estate returns. Real Estate Economics, 50(1), 89-121. https://doi.org/10.1111/1540-6229.12354", False),
    ("Ling, D. C., and Naranjo, A. (2015). Returns and information transmission dynamics in public and private "
     "real estate markets. Real Estate Economics, 43(1), 163-208. https://doi.org/10.1111/1540-6229.12069", False),
    ("", False),
    ("License", True),
    ("Data and documentation: Creative Commons Attribution-NonCommercial 4.0 International (CC BY-NC 4.0). "
     "Questions: open an issue at https://github.com/wang415705825/Real-Estate-Returns/issues", False),
]


def _sheet_from_csv(wb, title, path, int_cols=()):
    ws = wb.create_sheet(title)
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in df.itertuples(index=False):
        out = []
        for col, v in zip(df.columns, row):
            if v == "":
                out.append(None)
            elif col in int_cols:
                out.append(int(v))
            else:
                try:
                    out.append(float(v))
                except ValueError:
                    out.append(v)
        ws.append(out)
    for j, col in enumerate(df.columns, start=1):
        letter = ws.cell(row=1, column=j).column_letter
        ws.column_dimensions[letter].width = max(10, min(24, len(col) + 2))
        if col not in int_cols:
            for c in ws[letter][1:]:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
    ws.freeze_panes = "A2"


def workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws.column_dimensions["A"].width = 120
    for r, (text, bold) in enumerate(NOTES, start=1):
        c = ws.cell(row=r, column=1, value=text or None)
        c.font = Font(bold=bold, size=13 if r == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ints = {"year", "quarter", "n_all", "n_core", "n_noncore", "n_firms", "core", "n_quarters"}
    _sheet_from_csv(wb, "Indices", os.path.join(DATA, "reit_return_indices.csv"), ints)
    _sheet_from_csv(wb, "By property type", os.path.join(DATA, "reit_returns_by_property_type.csv"), ints)
    _sheet_from_csv(wb, "Summary statistics", os.path.join(DATA, "summary_statistics.csv"), ints)
    wb.properties.creator = "Real-Estate-Returns"
    wb.properties.title = "Unlevered and levered U.S. equity REIT return indices"
    path = os.path.join(DATA, "REIT_Return_Indices.xlsx")
    buf = io.BytesIO()
    wb.save(buf)
    _write_normalized_zip(buf.getvalue(), path)


def _write_normalized_zip(raw, path):
    """Rewrite an xlsx with fixed timestamps so identical content gives identical bytes."""
    src = zipfile.ZipFile(io.BytesIO(raw))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:)",
                              lambda m: m.group(1) + STAMP.encode() + m.group(3), data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            out.writestr(zi, data)


# ----------------------------------------------------------------------------- figures
def _use_fonts():
    paths = sorted(glob.glob(os.path.join(FONT_DIR, "*.ttf")))
    for p in paths:
        font_manager.fontManager.addfont(p)
    plt.rcParams.update({"font.family": "Inter" if paths else "DejaVu Sans",
                         "svg.hashsalt": "rer", "path.simplify": False})


def _figure(w, h, surface):
    return plt.figure(figsize=(w, h), dpi=200, facecolor=surface)


def _axes(fig, left, right, top, bottom):
    """Axes placed by margins in inches."""
    w, h = fig.get_size_inches()
    return fig.add_axes([left / w, bottom / h, 1 - (left + right) / w, 1 - (top + bottom) / h])


def _text(fig, x, y, s, **kw):
    """Figure text placed in inches from the top-left corner."""
    w, h = fig.get_size_inches()
    return fig.text(x / w, 1 - y / h, s, **kw)


def _style(ax, t, ygrid=True, xgrid=False):
    ax.set_facecolor(t["surface"])
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color(t["axis"])
    ax.spines["bottom"].set_linewidth(0.75)
    ax.tick_params(colors=t["muted"], labelsize=9, length=0, pad=6)
    ax.grid(False)
    if ygrid:
        ax.yaxis.grid(True, color=t["grid"], linewidth=0.75, linestyle="-")
    if xgrid:
        ax.xaxis.grid(True, color=t["grid"], linewidth=0.75, linestyle="-")
    ax.set_axisbelow(True)


def _header(fig, t, title, subtitle, size=15):
    _text(fig, 0.30, 0.26, title, ha="left", va="top", fontsize=size, fontweight="semibold", color=t["ink"])
    _text(fig, 0.30, 0.66, subtitle, ha="left", va="top", fontsize=9.5, color=t["ink2"])


def _source(fig, t, note=SOURCE):
    w, h = fig.get_size_inches()
    fig.text(0.30 / w, 0.16 / h, note, ha="left", va="bottom", fontsize=7.5, color=t["muted"])


def _key(fig, t, items, x, y):
    """One-row legend. items: (kind, color, label) with kind line, dot or patch; x, y in inches."""
    w, h = fig.get_size_inches()
    renderer = fig.canvas.get_renderer()
    yy = 1 - y / h
    for kind, color, label in items:
        if kind == "line":
            fig.add_artist(plt.Line2D([x / w, (x + 0.26) / w], [yy, yy], color=color, lw=1.8,
                                      solid_capstyle="round"))
            x += 0.36
        elif kind == "dot":
            fig.add_artist(plt.Line2D([(x + 0.06) / w], [yy], marker="o", ms=7, color=color, lw=0,
                                      mec=t["surface"], mew=1.2))
            x += 0.20
        else:
            fig.add_artist(Rectangle((x / w, yy - 0.06 / h), 0.26 / w, 0.12 / h, color=color, lw=0))
            x += 0.36
        txt = fig.text(x / w, yy, label, ha="left", va="center", fontsize=9, color=t["ink2"])
        x += txt.get_window_extent(renderer).width / fig.dpi + 0.30


def _save(fig, name, mode="light"):
    suffix = "-dark" if mode == "dark" else ""
    fig.savefig(os.path.join(FIGS, f"{name}{suffix}.png"), dpi=200, facecolor=fig.get_facecolor(),
                metadata={"Software": None})
    plt.close(fig)


def _growth(idx, col):
    """Growth of $1 from the end of 1992Q4, on a quarter-end year axis."""
    x = np.r_[1993.0, (idx.year + idx.quarter / 4.0).values]
    return x, np.r_[1.0, (1 + idx[col] / 100).cumprod().values]


def _pp(v):
    return f"{'+' if v >= 0 else '−'}{abs(v):.2f} pp"


def fig_cumulative(idx, mode):
    t = THEMES[mode]
    fig = _figure(8, 4.85, t["surface"])
    ax = _axes(fig, left=0.72, right=1.25, top=1.34, bottom=0.62)
    _style(ax, t)
    for a, b in RECESSIONS:
        ax.axvspan(a, b, color=t["band"], lw=0, zorder=0)
    x, unlev = _growth(idx, "vw_unlev_all")
    _, lev = _growth(idx, "vw_lev_all")
    wash = to_rgba(t["lev"], 0.13)
    ax.fill_between(x, unlev, lev, color=wash, lw=0, zorder=1)
    for key, level, name, side in [("lev", lev, "Levered", 1), ("unlev", unlev, "Unlevered", -1)]:
        ax.plot(x, level, color=t[key], lw=1.6, solid_joinstyle="round", solid_capstyle="round", zorder=3)
        ax.plot(x[-1], level[-1], "o", color=t[key], ms=6.5, mec=t["surface"], mew=1.5, zorder=4)
        ax.annotate(f"${level[-1]:.2f}", (x[-1], level[-1]), xytext=(9, 0), textcoords="offset points",
                    ha="left", va="center", fontsize=11, fontweight="semibold", color=t["ink"])
        ax.annotate(name, (x[-1], level[-1]), xytext=(9, 8 * side), textcoords="offset points",
                    ha="left", va="bottom" if side > 0 else "top", fontsize=8.5, color=t["ink2"])
    ax.set_yscale("log")
    ax.yaxis.set_major_locator(FixedLocator([1, 2, 5, 10, 20]))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:g}"))
    ax.yaxis.set_minor_locator(FixedLocator([]))
    ax.set_ylim(0.85, 28)
    ax.set_xlim(1992.6, 2026.2)
    ax.xaxis.set_major_locator(FixedLocator(list(range(1995, 2026, 5))))
    _header(fig, t, "REIT stocks vs. the assets behind them",
            "Growth of $1 invested at the start of 1993 in U.S. equity REITs, value-weighted; "
            "quarterly total returns, log scale")
    _key(fig, t, [("line", t["lev"], "Levered: REIT stock"), ("line", t["unlev"], "Unlevered: REIT assets"),
                  ("patch", wash, "Gap from leverage"), ("patch", t["band"], "U.S. recession (NBER)")],
         x=0.30, y=1.06)
    _source(fig, t)
    _save(fig, "cumulative_returns", mode)


def fig_property_types(sm, mode):
    t = THEMES[mode]
    full = sm[(sm.period_start == "1993Q1") & (sm.period_end == "2025Q4") & (sm.weighting == "VW")]
    groups = [g for g in full.group.unique() if g not in ("Core", "Non-core")]
    d = full[full.group.isin(groups)].pivot(index="group", columns="series", values="ann_return").round(2)
    d["gap"] = d.levered - d.unlevered                         # from the rounded returns, so labels add up
    types = d.drop("All").sort_values("gap")                  # bottom to top: largest gap on top
    d = pd.concat([types, d.loc[["All"]]])
    y = np.r_[np.arange(len(types)), len(types) + 0.35]         # All sits on top, set apart
    fig = _figure(8, 1.40 + 0.42 * (len(d) + 0.35) + 0.66, t["surface"])
    ax = _axes(fig, left=1.95, right=1.45, top=1.40, bottom=0.66)
    _style(ax, t, ygrid=False, xgrid=True)
    row = blended_transform_factory(fig.transFigure, ax.transData)
    w = fig.get_size_inches()[0]
    ax.add_patch(Rectangle((0.22 / w, y[-1] - 0.42), 1 - 0.44 / w, 0.84, transform=row, color=t["band"],
                           lw=0, zorder=0, clip_on=False))
    ax.hlines(y, d.unlevered, d.levered, color=t["muted"], lw=1.8, zorder=1)
    ax.scatter(d.unlevered, y, s=95, color=t["unlev"], edgecolors=t["surface"], linewidths=1.5, zorder=3)
    ax.scatter(d.levered, y, s=38, color=t["lev"], edgecolors=t["surface"], linewidths=1.5, zorder=3)
    labels = ["All equity REITs" if g == "All" else g for g in d.index]
    ax.set_yticks(y, labels, fontsize=9.5)
    ax.tick_params(axis="y", colors=t["ink2"], pad=10)
    top = ax.get_yticklabels()[-1]
    top.set_fontweight("semibold")
    top.set_color(t["ink"])
    lo, hi = sorted([d.loc["All", "unlevered"], d.loc["All", "levered"]])
    for v, off, ha in [(lo, -10, "right"), (hi, 10, "left")]:
        ax.annotate(f"{v:.2f}%", (v, y[-1]), xytext=(off, 0), textcoords="offset points",
                    ha=ha, va="center", fontsize=8.5, color=t["ink2"])
    xcol = 1 - 0.40 / w
    for yy, g, gap in zip(y, d.index, d.gap):
        ax.text(xcol, yy, _pp(gap), transform=row, ha="right", va="center", fontsize=9.5,
                fontweight="semibold" if g == "All" else "normal", color=t["ink"] if g == "All" else t["ink2"])
    ax.text(xcol, y[-1] + 0.95, "Levered − unlevered", transform=row, ha="right", va="center",
            fontsize=8.5, color=t["muted"])
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}%"))
    ax.xaxis.set_major_locator(MultipleLocator(2))
    ax.set_xlim(np.floor(d[["unlevered", "levered"]].min().min() - 0.6),
                np.ceil(d[["unlevered", "levered"]].max().max() + 0.6))
    ax.set_ylim(-0.6, y[-1] + 0.55)
    _header(fig, t, "What leverage added, by property type",
            "Annualized total return 1993–2025, value-weighted; property types with at least 3 REITs "
            "in every quarter")
    _key(fig, t, [("dot", t["unlev"], "Unlevered: REIT assets"), ("dot", t["lev"], "Levered: REIT stock")],
         x=0.30, y=1.06)
    _source(fig, t)
    _save(fig, "property_type_returns", mode)


def fig_ln2015(idx, mode):
    t = THEMES[mode]
    ws = load_workbook(LEGACY, read_only=True, data_only=True)["UnlevRet"]
    ln = {(int(r[0]), int(r[1])): float(r[2]) for r in ws.iter_rows(min_row=10, max_col=3, values_only=True)
          if r[0] is not None and r[2] is not None}
    ours = idx.set_index(["year", "quarter"]).vw_unlev_core
    m = pd.DataFrame({"ln": pd.Series(ln), "ours": ours}).dropna()
    fig = _figure(5.9, 6.6, t["surface"])
    ax = _axes(fig, left=0.95, right=0.35, top=1.05, bottom=0.95)
    _style(ax, t, ygrid=True, xgrid=True)
    lim = (-24, 18)
    ax.plot(lim, lim, color=t["muted"], linewidth=1.0, zorder=1)
    ax.scatter(m.ln, m.ours, s=36, color=t["unlev"], edgecolors=t["surface"], linewidths=1.2, zorder=3)
    ax.set_xlim(lim)
    ax.set_ylim(lim)
    ax.xaxis.set_major_locator(MultipleLocator(10))
    ax.yaxis.set_major_locator(MultipleLocator(10))
    ax.set_xlabel("Ling and Naranjo (2015), % per quarter", fontsize=9.5, color=t["ink2"], labelpad=8)
    ax.set_ylabel("This dataset, % per quarter", fontsize=9.5, color=t["ink2"], labelpad=6)
    first, last = m.index[0], m.index[-1]
    ax.text(0.04, 0.96, f"Correlation {m.ours.corr(m.ln):.3f}", transform=ax.transAxes, va="top",
            fontsize=11, fontweight="semibold", color=t["ink"])
    ax.text(0.04, 0.905, f"{len(m)} quarters, {first[0]}Q{first[1]}–{last[0]}Q{last[1]}\nGray line: equal returns",
            transform=ax.transAxes, va="top", fontsize=9, color=t["ink2"], linespacing=1.5)
    _header(fig, t, "Tracks Ling and Naranjo (2015) closely",
            "Value-weighted unlevered core REIT return, one dot per quarter")
    _source(fig, t, f"Data: Real-Estate-Returns ({REPO}).")
    _save(fig, "validation_ln2015", mode)


def fig_card(idx, sm, name, w, h):
    """Link-preview card (GitHub social preview, website and social posts); light only."""
    c = CARD
    full = sm[(sm.period_start == "1993Q1") & (sm.period_end == "2025Q4") & (sm.group == "All")
              & (sm.weighting == "VW")].set_index("series").ann_return
    fig = _figure(w, h, c["surface"])
    fig.add_artist(Rectangle((0, 1 - 0.07 / h), 1, 0.07 / h, color=c["accent"], lw=0))
    x0 = 0.42
    px, pw = 0.555 * w, 0.42 * w                                  # chart panel, inches
    py, ph = 0.40, h - 0.80
    _text(fig, x0, 0.38, "Open data  ·  1993–2025", ha="left", va="top", fontsize=8.5, fontweight="semibold",
          color=c["accent"])
    _text(fig, x0, 0.60, "U.S. REIT\nReturn Indices", ha="left", va="top", fontsize=23, fontweight="bold",
          color=c["ink"], linespacing=1.02)
    _text(fig, x0, 1.44, "Quarterly unlevered and levered total\nreturns for U.S. equity REITs",
          ha="left", va="top", fontsize=9.5, color=c["ink2"], linespacing=1.4)
    fig.add_artist(plt.Line2D([x0 / w, (x0 + 0.5) / w], [1 - 1.90 / h] * 2, color=c["rule"], lw=2))
    tiles = [(f"{full['levered']:.2f}%", "per year,\nREIT stock"),
             (f"{full['unlevered']:.2f}%", "per year,\nREIT assets"),
             (f"{idx.n_all.min()}–{idx.n_all.max()}", "REITs per\nquarter")]
    renderer = fig.canvas.get_renderer()
    xx = x0
    for value, label in tiles:
        a = _text(fig, xx, 2.02, value, ha="left", va="top", fontsize=15, fontweight="semibold", color=c["ink"])
        b = _text(fig, xx, 2.30, label, ha="left", va="top", fontsize=7, color=c["muted"], linespacing=1.3)
        xx += max(a.get_window_extent(renderer).width, b.get_window_extent(renderer).width) / fig.dpi + 0.30
    _text(fig, x0, h - 0.26, f"Chongyu Wang  ·  Florida State University  ·  {REPO}", ha="left",
          va="bottom", fontsize=7, color=c["ink2"])
    fig.add_artist(FancyBboxPatch((px / w, py / h), pw / w, ph / h, boxstyle="round,pad=0,rounding_size=0.012",
                                  transform=fig.transFigure, facecolor=c["panel"], edgecolor=c["axis"], lw=0.8,
                                  zorder=0))
    _text(fig, px + 0.18, h - py - ph + 0.18, "Growth of $1 since 1993", ha="left", va="top", fontsize=8.5,
          fontweight="semibold", color=c["ink2"])
    ax = fig.add_axes([(px + 0.42) / w, (py + 0.32) / h, (pw - 1.02) / w, (ph - 0.80) / h], zorder=1)
    t = dict(surface=c["panel"], axis=c["axis"], muted=c["muted"], grid=c["grid"])
    _style(ax, t)
    ax.tick_params(labelsize=7, pad=4)
    for key, col, label in [("lev", "vw_lev_all", "Levered"), ("unlev", "vw_unlev_all", "Unlevered")]:
        x, level = _growth(idx, col)
        ax.plot(x, level, color=c[key], lw=1.4, solid_joinstyle="round", solid_capstyle="round", zorder=3)
        ax.plot(x[-1], level[-1], "o", color=c[key], ms=5, mec=c["panel"], mew=1.2, zorder=4)
        ax.annotate(f"${level[-1]:.2f}", (x[-1], level[-1]), xytext=(6, 0), textcoords="offset points",
                    ha="left", va="center", fontsize=8.5, fontweight="semibold", color=c["ink"])
        ax.annotate(label, (x[-1], level[-1]), xytext=(6, 6 if key == "lev" else -6), textcoords="offset points",
                    ha="left", va="bottom" if key == "lev" else "top", fontsize=6.5, color=c["ink2"])
    ax.set_yscale("log")
    ax.yaxis.set_major_locator(FixedLocator([1, 5, 20]))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:g}"))
    ax.yaxis.set_minor_locator(FixedLocator([]))
    ax.set_ylim(0.85, 28)
    ax.set_xlim(1992.6, 2026.2)
    ax.xaxis.set_major_locator(FixedLocator([1995, 2005, 2015, 2025]))
    fig.savefig(os.path.join(FIGS, f"{name}.png"), dpi=200, facecolor=c["surface"], metadata={"Software": None})
    plt.close(fig)


def main():
    idx, bytype = load()
    sm = summary(idx, bytype)
    update_readme(readme_table(sm, idx))
    workbook()
    os.makedirs(FIGS, exist_ok=True)
    _use_fonts()
    for mode in THEMES:
        fig_cumulative(idx, mode)
        fig_ln2015(idx, mode)
        fig_property_types(sm, mode)
    fig_card(idx, sm, "social_preview", 6.4, 3.2)              # 1280 x 640, GitHub social preview
    fig_card(idx, sm, "social_card_1200x630", 6.0, 3.15)       # 1200 x 630, website and social posts
    print("wrote data/summary_statistics.csv, data/REIT_Return_Indices.xlsx, figures/*.png")


if __name__ == "__main__":
    main()
