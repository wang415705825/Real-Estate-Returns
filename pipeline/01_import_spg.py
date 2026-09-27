#!/usr/bin/env python3
"""
01_import_spg.py  —  Ingest the S&P Global Market Intelligence (formerly SNL)
"RE Companies" workbook into tidy long-format Stata files for the de-levering
pipeline.

Usage:
    python 01_import_spg.py --spg-file "RE Companies (S&P Global).xlsx" --data-dir /path/to/work

Input  : S&P Global MI "RE Companies" workbook export (licensed; not distributed)
         - 'Companies' sheet  : one row per REIT (classification + IDs)
         - 14 financial sheets: 7 vars x {Ann, Qtr}, WIDE (periods across columns)
Output : <data-dir>/spg_companies.dta, spg_fin_qtr.dta, spg_fin_ann.dta

Workbook quirks handled here:
  * junk rows: row1 'SNLTable', row3 field codes, row5 blank  -> data starts row6
  * period labels live in row4 (e.g. '2026Q1', '2025Y') for financial sheets
  * keyed across sheets by 'SNL Institution Key' (CUSIP/Ticker only in 'Companies')
  * UNIT MISMATCH: money vars in $000 EXCEPT Market Capitalization in $M.
      -> normalize every money series to $M (legacy WACC convention).
  * 'NA'/'NM' -> missing.
  * universe is GLOBAL incl. non-elected entities -> add us_cusip & elected flags
    (no hard filter here; downstream restricts via CRSP merge).
  * classification is point-in-time (one Property Type/Subtype per REIT).

Public copy of the production script of the same name; only path handling differs.
"""
import argparse
import os
import openpyxl
import pandas as pd

# financial sheet base name -> short Stata varname + unit tag
VARMAP = {
    "Total Assets":             ("TotalAssets",            "000"),
    "Total Debt":               ("TotalDebt",              "000"),
    "Total Preferred Equity":   ("TotalPreferredEquity",   "000"),
    "Market Capitalization":    ("MarketCapitalization",   "M"),
    "Interest Expense":         ("InterestExpense",        "000"),
    "Preferred Dividends Paid": ("PreferredDividendsPaid", "000"),
    "ROAE":                     ("ROAE",                   "pct"),
}
MISS = {"NA", "NM", "", None}


def _num(v, unit):
    """Coerce a cell to float in target units ($M for money; % as-is)."""
    if v in MISS:
        return None
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    if unit == "000":      # $000 -> $M
        return x / 1000.0
    return x               # $M and pct unchanged


def load_companies(wb):
    ws = wb["Companies"]
    rows = [r for r in ws.iter_rows(min_row=6, values_only=True) if r[0] is not None]
    df = pd.DataFrame(rows).iloc[:, :9]
    df.columns = ["name", "snlkey", "cusip", "ptype", "psub",
                  "elected", "reit_month", "reit_year", "ticker"]
    df["snlkey"] = pd.to_numeric(df["snlkey"], errors="coerce").astype("Int64")
    for c in ["cusip", "ticker", "name", "ptype", "psub", "elected"]:
        df[c] = df[c].astype("string").str.strip()
    for c in ["reit_month", "reit_year"]:
        df[c] = pd.to_numeric(
            df[c].replace({"NA": None}), errors="coerce").astype("Int64")
    df["elected_yes"] = (df["elected"] == "Yes").astype("int8")
    df["us_cusip"]    = df["cusip"].str[:1].str.isdigit().fillna(False).astype("int8")
    df["diversified"] = (df["ptype"] == "Diversified").astype("int8")
    df = df.drop_duplicates("snlkey").reset_index(drop=True)
    return df


def load_financial_sheet(wb, sheet, varname, unit):
    ws = wb[sheet]
    periods = [c for c in next(ws.iter_rows(min_row=4, max_row=4, values_only=True))]
    data = [r for r in ws.iter_rows(min_row=6, values_only=True) if r[1] is not None]
    recs = []
    for r in data:
        snlkey = r[1]
        for j in range(3, len(periods)):          # cols D.. are period values
            per = periods[j]
            if per in (None, ""):
                continue
            val = _num(r[j] if j < len(r) else None, unit)
            if val is None:
                continue
            recs.append((snlkey, str(per).strip(), val))
    out = pd.DataFrame(recs, columns=["snlkey", "period", varname])
    out["snlkey"] = pd.to_numeric(out["snlkey"], errors="coerce").astype("Int64")
    return out


def build_panel(wb, freq):
    suffix = "(Ann)" if freq == "ann" else "(Qtr)"
    merged = None
    for base, (var, unit) in VARMAP.items():
        sheet = f"{base} {suffix}"
        s = load_financial_sheet(wb, sheet, var, unit)
        merged = s if merged is None else merged.merge(s, on=["snlkey", "period"], how="outer")
    # parse period -> year (+ quarter)
    if freq == "ann":
        merged["year"] = merged["period"].str.replace("Y", "", regex=False).astype(int)
    else:
        yq = merged["period"].str.extract(r"(\d{4})Q([1-4])")
        merged["year"]    = yq[0].astype(int)
        merged["quarter"] = yq[1].astype(int)
    return merged


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--spg-file", required=True,
                    help='S&P Global MI "RE Companies" workbook (.xlsx)')
    ap.add_argument("--data-dir", required=True,
                    help="working directory for licensed intermediate files "
                         "(keep it outside the repository)")
    args = ap.parse_args()
    INFILE = args.spg_file
    OUTDIR = args.data_dir

    wb = openpyxl.load_workbook(INFILE, read_only=True, data_only=True)

    comp = load_companies(wb)
    qtr  = build_panel(wb, "qtr")
    ann  = build_panel(wb, "ann")

    # attach identifiers/classification from Companies
    keys = comp[["snlkey", "cusip", "ticker", "ptype", "psub",
                 "elected_yes", "us_cusip", "diversified"]]
    qtr = qtr.merge(keys, on="snlkey", how="left")
    ann = ann.merge(keys, on="snlkey", how="left")

    # tidy column order
    idv = ["snlkey", "cusip", "ticker", "ptype", "psub",
           "elected_yes", "us_cusip", "diversified"]
    finv = list(v for v, _ in VARMAP.values())
    qtr = qtr[idv + ["year", "quarter", "period"] + finv].sort_values(["snlkey", "year", "quarter"])
    ann = ann[idv + ["year", "period"] + finv].sort_values(["snlkey", "year"])

    os.makedirs(OUTDIR, exist_ok=True)
    comp.to_stata(os.path.join(OUTDIR, "spg_companies.dta"), write_index=False, version=118)
    qtr.to_stata(os.path.join(OUTDIR, "spg_fin_qtr.dta"),   write_index=False, version=118)
    ann.to_stata(os.path.join(OUTDIR, "spg_fin_ann.dta"),   write_index=False, version=118)
    wb.close()

    # ---- report ----
    print("WROTE ->", OUTDIR)
    print(f"  spg_companies.dta : {len(comp):>6} REITs "
          f"(US&elected={((comp.us_cusip==1)&(comp.elected_yes==1)).sum()}, "
          f"diversified={comp.diversified.sum()})")
    print(f"  spg_fin_qtr.dta   : {len(qtr):>6} rows | "
          f"{qtr.year.min()}Q? .. {qtr.year.max()}Q? | "
          f"{qtr.snlkey.nunique()} REITs w/ qtr data")
    print(f"  spg_fin_ann.dta   : {len(ann):>6} rows | "
          f"{ann.year.min()} .. {ann.year.max()} | "
          f"{ann.snlkey.nunique()} REITs w/ ann data")
    # sanity: unit normalization spot-check (Total Assets should be in $M now)
    chk = qtr.dropna(subset=["TotalAssets"]).sort_values("year").tail(1)
    if len(chk):
        r = chk.iloc[0]
        print(f"  unit check: {r.cusip} {int(r.year)}Q{int(r.quarter)} "
              f"TotalAssets={r.TotalAssets:.1f}$M, MktCap={r.MarketCapitalization}$M")


if __name__ == "__main__":
    main()
